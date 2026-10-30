"""
Turns a student's Word / Excel / PowerPoint file into a plain-text inspection report for the grading model.

Plain text extraction loses exactly what Office homework is about (bold, sizes, alignment, lists, tables, formulas,
charts, slide layouts), so the model could not tell a student who applied a command from one who only described it.
The report lists what is really in the file, element by element, in document order.
"""
import io
import zipfile
from typing import Optional

MAX_REPORT_CHARS = 14000
MAX_SHEETS = 6
MAX_CELLS_PER_SHEET = 400
MAX_ROWS_SCANNED = 300
MAX_COLS_SCANNED = 40
MAX_SLIDES = 40


def office_kind(data: bytes) -> Optional[str]:
    if data[:2] != b"PK":
        return None
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            names = set(z.namelist())
    except zipfile.BadZipFile:
        return None
    if "word/document.xml" in names:
        return "docx"
    if "xl/workbook.xml" in names:
        return "xlsx"
    if "ppt/presentation.xml" in names:
        return "pptx"
    return None


def describe_office(data: bytes) -> str:
    kind = office_kind(data)
    try:
        if kind == "docx":
            report = _describe_docx(data)
        elif kind == "xlsx":
            report = _describe_xlsx(data)
        elif kind == "pptx":
            report = _describe_pptx(data)
        else:
            return ""
    except Exception:
        return ""
    if len(report) > MAX_REPORT_CHARS:
        report = report[:MAX_REPORT_CHARS] + "\n[report cut: the file is longer]"
    return report


def _quote(text: str, limit: int = 300) -> str:
    text = " ".join((text or "").split())
    return "«" + (text[:limit] + "…" if len(text) > limit else text) + "»"


def _color(value) -> Optional[str]:
    try:
        rgb = str(value) if value is not None else None
    except Exception:
        return None
    if not rgb or rgb in ("None", "000000", "FF000000", "00000000"):
        return None
    return "#" + rgb[-6:]


def _describe_docx(data: bytes) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(data))

    def style_value(style, getter):
        while style is not None:
            value = getter(style)
            if value is not None:
                return value
            style = style.base_style
        return None

    def run_format(run, paragraph):
        font, style = run.font, paragraph.style

        def pick(attr):
            value = getattr(font, attr)
            if value is None and run.style is not None:
                value = style_value(run.style, lambda s: getattr(s.font, attr))
            if value is None:
                value = style_value(style, lambda s: getattr(s.font, attr))
            return value

        marks = []
        if pick("bold"):
            marks.append("bold")
        if pick("italic"):
            marks.append("italic")
        if pick("underline"):
            marks.append("underline")
        if pick("strike"):
            marks.append("strikethrough")
        if pick("superscript"):
            marks.append("superscript")
        if pick("subscript"):
            marks.append("subscript")
        size = pick("size")
        if size is not None:
            marks.append(f"{size.pt:g}pt")
        name = pick("name")
        if name:
            marks.append(f"font {name}")
        try:
            color = _color(font.color.rgb) if font.color is not None and font.color.type is not None else None
        except Exception:
            color = None
        if color:
            marks.append(f"color {color}")
        if font.highlight_color is not None:
            marks.append(f"highlight {str(font.highlight_color).split('.')[-1].split(' ')[0].lower()}")
        return ", ".join(marks)

    def paragraph_line(paragraph: Paragraph, label: str) -> Optional[str]:
        segments, last_format, has_picture = [], None, False
        for run in paragraph.runs:
            if run._element.xpath(".//*[local-name()='drawing' or local-name()='pict']"):
                has_picture = True
            if not run.text:
                continue
            fmt = run_format(run, paragraph)
            if segments and fmt == last_format:
                segments[-1][0] += run.text
            else:
                segments.append([run.text, fmt])
                last_format = fmt
        text = "".join(s[0] for s in segments).strip()
        if not text and not has_picture:
            return None
        props = [f"style={paragraph.style.name}"] if paragraph.style is not None else []
        alignment = paragraph.alignment
        if alignment is None:
            alignment = style_value(paragraph.style, lambda s: s.paragraph_format.alignment)
        if alignment is not None:
            props.append("align=" + str(alignment).split(".")[-1].split(" ")[0].lower())
        num = paragraph._p.pPr.numPr if paragraph._p.pPr is not None else None
        if num is not None or (paragraph.style is not None and "list" in paragraph.style.name.lower()):
            level = num.ilvl.val if num is not None and num.ilvl is not None else 0
            props.append(f"list item level {level + 1}")
        if paragraph._p.xpath(".//*[local-name()='br'][@*[local-name()='type']='page']"):
            props.append("page break")
        line = f"{label} [{', '.join(props)}]:"
        if has_picture:
            line += " [picture]"
        if len(segments) == 1:
            line += f" {_quote(segments[0][0])}" + (f" {{{segments[0][1]}}}" if segments[0][1] else "")
        else:
            for chunk, fmt in segments:
                if chunk.strip():
                    line += f" {_quote(chunk, 160)}" + (f" {{{fmt}}}" if fmt else "")
        return line

    lines = ["[Word document inspection: what is really in the file, in order. {..} = formatting applied to that text]"]
    section = doc.sections[0] if doc.sections else None
    if section is not None and section.page_width and section.page_height:
        orientation = "landscape" if section.page_width > section.page_height else "portrait"
        margins = [m.cm for m in (section.top_margin, section.bottom_margin, section.left_margin, section.right_margin) if m is not None]
        lines.append(f"Page: {orientation}, {section.page_width.cm:.1f} x {section.page_height.cm:.1f} cm"
                     + (f", margins top/bottom/left/right {' / '.join(f'{m:.1f}' for m in margins)} cm" if margins else ""))
        columns = section._sectPr.xpath("./*[local-name()='cols']/@*[local-name()='num']")
        if columns and columns[0] not in ("1", "0"):
            lines.append(f"Text columns: {columns[0]}")
        header = " ".join(p.text for p in section.header.paragraphs if p.text.strip()) if not section.header.is_linked_to_previous else ""
        footer = " ".join(p.text for p in section.footer.paragraphs if p.text.strip()) if not section.footer.is_linked_to_previous else ""
        if header.strip():
            lines.append(f"Header: {_quote(header)}")
        if footer.strip():
            lines.append(f"Footer: {_quote(footer)}")
        if section._sectPr.xpath(".//*[local-name()='pgNumType']") or "PAGE" in (section.footer._element.xml + section.header._element.xml):
            lines.append("Page numbers: yes")

    paragraph_no = table_no = 0
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph_no += 1
            line = paragraph_line(Paragraph(child, doc), f"Paragraph {paragraph_no}")
            if line:
                lines.append(line)
        elif tag == "tbl":
            table_no += 1
            table = Table(child, doc)
            rows = len(table.rows)
            cols = len(table.columns)
            style = table.style.name if table.style is not None else "none"
            borders = bool(child.xpath(".//*[local-name()='tblBorders' or local-name()='tcBorders']")) or style != "none"
            lines.append(f"Table {table_no}: {rows} rows x {cols} columns, style={style}, borders={'yes' if borders else 'no'}")
            for r, row in enumerate(table.rows[:15], start=1):
                cells = []
                for cell in row.cells[:10]:
                    text = " ".join(cell.text.split())
                    bold = any(run.bold for p in cell.paragraphs for run in p.runs if run.text.strip())
                    cells.append((text[:60] or "(empty)") + (" {bold}" if bold else ""))
                lines.append(f"  row {r}: " + " | ".join(cells))
            if rows > 15:
                lines.append(f"  ... {rows - 15} more rows")
        if sum(len(x) for x in lines) > MAX_REPORT_CHARS:
            break

    pictures = len(doc.inline_shapes)
    floating = len(doc.element.body.xpath(".//*[local-name()='anchor']"))
    if pictures or floating:
        lines.append(f"Pictures/shapes in the document: {pictures} inline, {floating} floating")
    if doc.element.body.xpath(".//*[local-name()='hyperlink']"):
        lines.append("Hyperlinks: yes")
    if doc.element.body.xpath(".//*[local-name()='txbxContent']"):
        boxes = [" ".join(t for t in box.itertext()) for box in doc.element.body.xpath(".//*[local-name()='txbxContent']")]
        lines.append("Text boxes: " + " ; ".join(_quote(b, 120) for b in boxes[:5]))
    if doc.element.body.xpath(".//*[local-name()='fldSimple' and contains(@*[local-name()='instr'],'TOC')]") or \
            "TOC \\" in doc.element.body.xml:
        lines.append("Table of contents: yes")
    return "\n".join(lines)


def _chart_label(chart) -> str:
    label = type(chart).__name__.replace("Chart", "").lower() + " chart"
    try:
        rich = chart.title.tx.rich if chart.title is not None and chart.title.tx is not None else None
        title = "".join(run.t for paragraph in rich.p for run in (paragraph.r or [])) if rich is not None else ""
    except Exception:
        title = ""
    return label + (f" titled {_quote(title, 80)}" if title.strip() else "")


def _describe_xlsx(data: bytes) -> str:
    import openpyxl
    from openpyxl.utils import get_column_letter

    formulas = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    values = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    lines = ["[Excel workbook inspection: what is really in the file. '=...' is a formula typed in the cell, "
             "'-> v' is its saved result; a number without '=' was typed by hand. {..} = cell formatting]"]
    lines.append(f"Sheets: {', '.join(formulas.sheetnames)}")
    for chartsheet in formulas.chartsheets:
        kinds = ", ".join(_chart_label(c) for c in getattr(chartsheet, "_charts", []))
        lines.append(f"Chart sheet {_quote(chartsheet.title, 60)} (a chart on its own sheet)" + (f": {kinds}" if kinds else ""))
    for ws in formulas.worksheets[:MAX_SHEETS]:
        wv = values[ws.title]
        lines.append(f"Sheet {_quote(ws.title, 60)}: used range {ws.dimensions}")
        centered = [side for side, on in (("horizontally", ws.print_options.horizontalCentered),
                                          ("vertically", ws.print_options.verticalCentered)) if on]
        if centered:
            lines.append(f"  page setup: centered on the page {' and '.join(centered)}")
        if ws.page_setup is not None and ws.page_setup.orientation:
            lines.append(f"  page orientation: {ws.page_setup.orientation}")
        widths = [f"{letter}={dim.width:g}" for letter, dim in ws.column_dimensions.items()
                  if dim.width and (dim.customWidth or dim.bestFit)]
        if widths:
            lines.append("  custom column widths: " + ", ".join(widths[:30]))
        if ws.freeze_panes:
            lines.append(f"  frozen panes at {ws.freeze_panes}")
        if ws.auto_filter and ws.auto_filter.ref:
            lines.append(f"  filter on {ws.auto_filter.ref}")
        if ws.merged_cells.ranges:
            lines.append("  merged cells: " + ", ".join(str(r) for r in list(ws.merged_cells.ranges)[:20]))
        if ws.tables:
            lines.append("  Excel tables: " + ", ".join(f"{name} ({t.ref})" for name, t in ws.tables.items()))
        rules = sum(len(r.rules) for r in ws.conditional_formatting)
        if rules:
            lines.append(f"  conditional formatting rules: {rules} on " + ", ".join(str(r.sqref) for r in ws.conditional_formatting)[:200])
        if ws.data_validations and ws.data_validations.dataValidation:
            lines.append("  data validation on: " + ", ".join(str(v.sqref) for v in ws.data_validations.dataValidation)[:200])
        for chart in getattr(ws, "_charts", []):
            lines.append(f"  chart: {_chart_label(chart)}")
        if getattr(ws, "_images", []):
            lines.append(f"  pictures: {len(ws._images)}")

        shown = 0
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, MAX_ROWS_SCANNED), max_col=min(ws.max_column, MAX_COLS_SCANNED)):
            for cell in row:
                if cell.value is None:
                    continue
                ref = f"{get_column_letter(cell.column)}{cell.row}"
                raw = cell.value
                if isinstance(raw, str) and raw.startswith("="):
                    result = wv[ref].value
                    text = f"{raw}" + (f" -> {result}" if result is not None else " -> (no saved result)")
                else:
                    text = _quote(raw, 80) if isinstance(raw, str) else str(raw)
                marks = []
                if cell.font is not None:
                    if cell.font.b:
                        marks.append("bold")
                    if cell.font.i:
                        marks.append("italic")
                    if cell.font.u:
                        marks.append("underline")
                    if cell.font.sz and float(cell.font.sz) != 11:
                        marks.append(f"{float(cell.font.sz):g}pt")
                    color = _color(cell.font.color.rgb) if cell.font.color is not None and cell.font.color.type == "rgb" else None
                    if color:
                        marks.append(f"text {color}")
                fill = cell.fill
                if fill is not None and fill.fill_type == "solid" and fill.fgColor is not None and fill.fgColor.type == "rgb":
                    color = _color(fill.fgColor.rgb)
                    if color:
                        marks.append(f"fill {color}")
                if cell.number_format and cell.number_format != "General":
                    marks.append(f"format {cell.number_format}")
                if cell.alignment is not None and cell.alignment.horizontal:
                    marks.append(f"align {cell.alignment.horizontal}")
                if cell.border is not None and any(getattr(cell.border, side).style for side in ("left", "right", "top", "bottom")):
                    marks.append("border")
                lines.append(f"  {ref} = {text}" + (f" {{{', '.join(marks)}}}" if marks else ""))
                shown += 1
                if shown >= MAX_CELLS_PER_SHEET:
                    break
            if shown >= MAX_CELLS_PER_SHEET:
                lines.append("  ... more cells")
                break
    if len(formulas.worksheets) > MAX_SHEETS:
        lines.append(f"... {len(formulas.worksheets) - MAX_SHEETS} more sheets")
    return "\n".join(lines)


def _describe_pptx(data: bytes) -> str:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    pres = Presentation(io.BytesIO(data))
    lines = ["[PowerPoint inspection: what is really on each slide, in order. {..} = text formatting]"]
    lines.append(f"Slides: {len(pres.slides)}")

    def text_lines(frame, indent):
        out = []
        for paragraph in frame.paragraphs:
            text = "".join(r.text for r in paragraph.runs).strip()
            if not text:
                continue
            marks = []
            runs = [r for r in paragraph.runs if r.text.strip()]
            if runs and all(r.font.bold for r in runs):
                marks.append("bold")
            if runs and all(r.font.italic for r in runs):
                marks.append("italic")
            sizes = {r.font.size.pt for r in runs if r.font.size is not None}
            if sizes:
                marks.append("/".join(f"{s:g}pt" for s in sorted(sizes)))
            names = {r.font.name for r in runs if r.font.name}
            if names:
                marks.append("font " + "/".join(sorted(names)))
            level = f"level {paragraph.level + 1} " if paragraph.level else ""
            out.append(f"{indent}{level}{_quote(text, 200)}" + (f" {{{', '.join(marks)}}}" if marks else ""))
        return out

    def shape_lines(shape, indent):
        out = []
        kind = shape.shape_type
        if kind == MSO_SHAPE_TYPE.GROUP:
            out.append(f"{indent}group:")
            for inner in shape.shapes:
                out += shape_lines(inner, indent + "  ")
            return out
        if kind == MSO_SHAPE_TYPE.PICTURE:
            out.append(f"{indent}picture")
            return out
        if getattr(shape, "has_chart", False) and shape.has_chart:
            chart = shape.chart
            title = chart.chart_title.text_frame.text if chart.has_title else ""
            out.append(f"{indent}chart: {str(chart.chart_type).split('.')[-1].split(' ')[0].lower()}" + (f" titled {_quote(title, 80)}" if title else ""))
            return out
        if getattr(shape, "has_table", False) and shape.has_table:
            table = shape.table
            out.append(f"{indent}table {len(table.rows)} rows x {len(table.columns)} columns")
            for r, row in enumerate(list(table.rows)[:10], start=1):
                out.append(f"{indent}  row {r}: " + " | ".join((" ".join(c.text.split())[:50] or "(empty)") for c in row.cells))
            return out
        if shape.is_placeholder:
            label = f"placeholder {str(shape.placeholder_format.type).split('.')[-1].split(' ')[0].lower()}"
        elif kind == MSO_SHAPE_TYPE.AUTO_SHAPE:
            try:
                label = f"shape {str(shape.auto_shape_type).split('.')[-1].split(' ')[0].lower()}"
            except Exception:
                label = "shape"
        elif kind == MSO_SHAPE_TYPE.TEXT_BOX:
            label = "text box"
        elif kind == MSO_SHAPE_TYPE.MEDIA:
            label = "video/audio"
        else:
            label = str(kind).split(".")[-1].split(" ")[0].lower() if kind is not None else "object"
        texts = text_lines(shape.text_frame, indent + "  ") if getattr(shape, "has_text_frame", False) and shape.has_text_frame else []
        if texts or label != "object":
            out.append(f"{indent}{label}:" if texts else f"{indent}{label}")
            out += texts
        return out

    for number, slide in enumerate(list(pres.slides)[:MAX_SLIDES], start=1):
        xml = slide._element.xml
        extras = []
        if "<p:transition" in xml:
            extras.append("transition")
        if "<p:timing" in xml:
            extras.append("animations")
        if slide.background is not None and slide._element.cSld.bg is not None:
            extras.append("custom background")
        lines.append(f"Slide {number} (layout {_quote(slide.slide_layout.name, 40)}" + (", " + ", ".join(extras) if extras else "") + "):")
        for shape in slide.shapes:
            lines += shape_lines(shape, "  ")
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"  speaker notes: {_quote(notes, 200)}")
    if len(pres.slides) > MAX_SLIDES:
        lines.append(f"... {len(pres.slides) - MAX_SLIDES} more slides")
    return "\n".join(lines)
